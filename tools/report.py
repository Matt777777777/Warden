#!/usr/bin/env python3
"""
WARDEN - inspection report drafted by a local language model, checked by code.

    python tools/report.py                  # reads the crate's log over USB
    python tools/report.py --file dump.txt  # a dump saved by verify_log.py --save
    python tools/report.py --demo           # a built-in sample shipment, no hardware needed
    python tools/report.py --model llama3.2:1b   # smaller, faster model

Pipeline: the log is verified (hash chain), the facts are extracted and the risk level
is decided by plain code. Only then does a small language model, running offline on
this laptop through Ollama, write the prose: summary, what each incident probably
means, what the inspector should do. Every time and figure it writes is checked
against the verified log; a sentence that cannot be matched is removed. The model is
the clerk who writes the report, not the investigator who decides what happened.

Output: reports/<crate>_<date>.docx and the same report as JSON (used by the phone app).
Requires Ollama (https://ollama.com) with a model pulled, e.g. `ollama pull llama3.2:3b`.
Without it, a plain template is used and the report says so.
"""

import argparse
import calendar
import json
import re
import sys
import time
import urllib.error
import urllib.request
from datetime import datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import verify_log as V  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
REPORTS = ROOT / "reports"
DEFAULT_MODEL = "llama3.2:3b"
OLLAMA = "http://localhost:11434"

LABELS = {
    "BOOT": "Power on", "OPEN OK": "Authorised opening", "OPEN !!": "Lid opened without authorisation",
    "CLOSED": "Lid closed", "DROP": "Drop", "SHOCK": "Shock (firmware v0.3)", "FLIPPED": "Turned upside down",
    "UPRIGHT": "Back upright", "AUTH": "Authorised badge", "INSPECT": "Inspector badge",
    "BADGE ?": "Unknown badge", "CLEARED": "Log cleared", "TAMPER": "Forced-entry attempt (on-device AI)",
}


# ---------------------------------------------------------------- facts (code, not the model)
def hhmm(r):
    return r["dt"][11:16]


def detail(r):
    if r["name"] in ("DROP", "SHOCK"):
        return f"impact {r['value'] / 10:.1f} g"
    if r["name"] == "TAMPER":
        return f"{r['value']}/10 trees agreed"
    if r["name"] == "BADGE ?":
        return f"badge ID starting {r['value']:02X}"
    return ""


def minutes_between(a, b):
    return round((b["unix"] - a["unix"]) / 60)


def build_facts(chk, crate_id):
    recs = chk["records"]
    facts = {"crate_id": crate_id, "integrity_ok": chk["ok"], "integrity_issues": chk["issues"],
             "records": len(recs), "incidents": [], "authorised_openings": [], "counts": {}}
    if not recs:
        facts.update(start=None, end=None)
        return facts
    facts["start"], facts["end"] = recs[0]["dt"], recs[-1]["dt"]
    for r in recs:
        facts["counts"][LABELS.get(r["name"], r["name"])] = facts["counts"].get(LABELS.get(r["name"], r["name"]), 0) + 1

    def next_event(i, names, within_min=None):
        for r in recs[i + 1:]:
            if within_min is not None and r["unix"] - recs[i]["unix"] > within_min * 60:
                return None
            if r["name"] in names:
                return r
        return None

    def prev_event(i, names, within_min):
        for r in reversed(recs[:i]):
            if recs[i]["unix"] - r["unix"] > within_min * 60:
                return None
            if r["name"] in names:
                return r
        return None

    for i, r in enumerate(recs):
        name, ctx, severity = r["name"], [], None
        if name == "OPEN OK":
            closed = next_event(i, {"CLOSED"})
            facts["authorised_openings"].append(
                f"{r['dt']}" + (f", lid open {minutes_between(r, closed)} min" if closed else ", not closed since"))
            continue
        if name == "OPEN !!":
            severity = "high"
            t = prev_event(i, {"TAMPER"}, 5)
            if t:
                ctx.append(f"preceded by a forced-entry signature at {hhmm(t)}")
            closed = next_event(i, {"CLOSED"})
            ctx.append(f"lid open {minutes_between(r, closed)} min, closed at {hhmm(closed)}" if closed
                       else "lid not closed since")
        elif name == "TAMPER":
            severity = "high"
            o = next_event(i, {"OPEN !!", "OPEN OK"}, 5)
            ctx.append(f"lid opened at {hhmm(o)}" if o else "lid stayed closed afterwards")
        elif name == "BADGE ?":
            severity = "high"
            ctx.append("badge not on the crate's list")
        elif name == "DROP":
            severity = "medium"
            ctx.append("free fall followed by impact")
        elif name == "SHOCK":
            severity = "medium"
        elif name == "FLIPPED":
            severity = "medium"
            up = next_event(i, {"UPRIGHT"})
            ctx.append(f"upside down {minutes_between(r, up)} min, upright at {hhmm(up)}" if up
                       else "not turned back upright since")
        elif name == "BOOT" and i > 0:
            severity = "medium"
            ctx.append(f"crate restarted; last event before it at {hhmm(recs[i - 1])}: power was cut at some point "
                       "in between and nothing was monitored while it was off")
        elif name == "CLEARED":
            severity = "high"
            ctx.append("the log was erased through the maintenance port; earlier history is gone")
        if severity:
            facts["incidents"].append({"id": f"I{len(facts['incidents']) + 1}", "time": r["dt"], "event": name,
                                       "label": LABELS.get(name, name), "detail": detail(r),
                                       "context": "; ".join(ctx), "severity": severity})
    facts["type_counts"] = {}
    for x in facts["incidents"]:
        facts["type_counts"][x["event"]] = facts["type_counts"].get(x["event"], 0) + 1
    facts["sequences"] = linked_sequences(recs)
    return facts


SEQ_WORDS = {"TAMPER": "forced-entry signature", "OPEN !!": "lid opened without authorisation",
             "CLOSED": "lid closed", "BOOT": "crate restarted after a power cut", "BADGE ?": "unknown badge",
             "CLEARED": "log erased", "OPEN OK": "authorised opening", "DROP": "drop", "FLIPPED": "turned upside down"}
SECURITY = {"TAMPER", "OPEN !!", "BOOT", "BADGE ?", "CLEARED"}


def linked_sequences(recs, gap_min=60):
    """Security events less than an hour apart, found by code so the model does not have to infer them."""
    seqs, cur = [], []
    for i, r in enumerate(recs):
        security = r["name"] in SECURITY and not (r["name"] == "BOOT" and i == 0)
        if cur and r["unix"] - cur[-1]["unix"] > gap_min * 60:
            if sum(x["name"] in SECURITY for x in cur) >= 2:
                seqs.append(cur)
            cur = []
        if security or (cur and r["name"] in ("CLOSED", "OPEN OK")):
            cur.append(r)
    if sum(x["name"] in SECURITY for x in cur) >= 2:
        seqs.append(cur)
    return [", then ".join(f"{hhmm(r)} {SEQ_WORDS.get(r['name'], r['name'])}" for r in sq)
            + f" ({minutes_between(sq[0], sq[-1])} min in total)" for sq in seqs]


NUMBERS = ["no", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten"]


def count_phrase(n, singular, plural):
    """'one unauthorised lid opening', 'two drops': small models misread '(1x)' as 'several'."""
    return f"{NUMBERS[n] if n < len(NUMBERS) else n} {singular if n == 1 else plural}"


def risk_level(facts):
    """Decided by rules, never by the language model."""
    reasons = []
    if not facts["integrity_ok"]:
        reasons.append("the log failed its integrity check")
    for kind, one, many in (("OPEN !!", "unauthorised lid opening", "unauthorised lid openings"),
                            ("TAMPER", "forced-entry attempt", "forced-entry attempts"),
                            ("BADGE ?", "unknown badge presented", "unknown badges presented"),
                            ("CLEARED", "log erasure during the shipment", "log erasures during the shipment")):
        n = facts["type_counts"].get(kind, 0)
        if n:
            reasons.append(count_phrase(n, one, many))
    if reasons:
        return "HIGH", reasons
    for kind, one, many in (("DROP", "drop", "drops"), ("FLIPPED", "time turned upside down", "times turned upside down"),
                            ("BOOT", "power interruption", "power interruptions"), ("SHOCK", "shock", "shocks")):
        n = facts["type_counts"].get(kind, 0)
        if n:
            reasons.append(count_phrase(n, one, many))
    if reasons:
        return "MEDIUM", reasons
    return "LOW", ["no incident recorded"]


def facts_text(facts, level, reasons):
    lines = [f"Crate: {facts['crate_id']}",
             f"Log period (crate clock): {facts['start']} to {facts['end']}",
             f"Log integrity: {'VERIFIED, hash chain unbroken' if facts['integrity_ok'] else 'FAILED: ' + '; '.join(facts['integrity_issues'])}",
             f"Records: {facts['records']}",
             f"Overall risk level (decided by rules): {level}, because: {'; '.join(reasons)}",
             f"Authorised openings: {len(facts['authorised_openings'])}"
             + (" (" + "; ".join(facts["authorised_openings"]) + ")" if facts["authorised_openings"] else ""),
             "Incidents:"]
    if not facts["incidents"]:
        lines.append("  none")
    for x in facts["incidents"][:15]:
        lines.append(f"  {x['id']}: {x['time']} {x['label']}"
                     + (f", {x['detail']}" if x["detail"] else "") + (f" ({x['context']})" if x["context"] else ""))
    if facts["sequences"]:
        lines.append("Linked sequences (security events less than an hour apart):")
        lines += [f"  {sq}" for sq in facts["sequences"]]
    return "\n".join(lines)


# ---------------------------------------------------------------- the language model (prose only)
SYSTEM = """You draft end-use monitoring inspection reports for military equipment crates.
You receive verified facts from the crate's tamper-evident log. Write for a customs or arms-control inspector, in plain, sober English.
Rules: use ONLY the facts given. Never invent times, places, people, numbers or causes. Quote times exactly as written (24-hour clock).
State likely explanations as possibilities ("may indicate"), never as certainties. Do not change the risk level you are given.
Respect the counts: if something happened once, say once, never "multiple" or "repeated".

What the events mean:
- Forced-entry attempt: the crate's on-board motion classifier recognised someone prying, hammering or forcing the closed lid.
- Lid opened without authorisation: light inside the crate while no authorised badge had been presented in the previous 30 seconds (no badge at all was used, not a wrong one). After such an opening the contents must be counted against the packing list.
- Power on, after the first record: the crate restarted, so its power was cut; nothing is recorded while it is off. A power cut right after an intrusion may be an attempt to hide what happened next.
- Unknown badge: a badge not registered for this crate was presented, possibly to try to obtain an authorised opening.
- Drop: free fall then impact; the contents may be damaged. Turned upside down: the contents may have shifted.
- A linked sequence is the strongest evidence: describe it as one event, in order, quoting its times.
- With a HIGH risk level, never write that no investigation or action is needed. If the log integrity is VERIFIED, do not suggest the log itself was manipulated.
- Refer to incidents by their time, never by their id (I1, I2...).

Style: specific, factual, short. Each assessment must use that incident's own facts (time, force, duration, what came before or after) and must not repeat another assessment. Never write "further investigation is required"; say what exactly to check."""

TASK = """Facts:
{facts}

Write a JSON object with:
- "executive_summary": 3 to 4 sentences for the inspector: the risk level, the most serious incident or linked sequence with its times, and whether the log can be trusted.
- "assessments": one entry per incident, {{"id": incident id, "assessment": 1 to 2 sentences on what this specific incident may indicate and what exactly to check}}.
- "recommendations": 3 concrete actions for the inspector, each tied to a specific incident or time window."""

SCHEMA = {
    "type": "object",
    "properties": {
        "executive_summary": {"type": "string"},
        "assessments": {"type": "array", "items": {"type": "object", "properties": {
            "id": {"type": "string"}, "assessment": {"type": "string"}}, "required": ["id", "assessment"]}},
        "recommendations": {"type": "array", "items": {"type": "string"}, "minItems": 3, "maxItems": 4},
    },
    "required": ["executive_summary", "assessments", "recommendations"],
}


def ask_llm(prompt, model, progress=None, host=OLLAMA):
    """Streams the answer from Ollama so the app can show progress. Returns (dict, stats) or raises."""
    body = json.dumps({"model": model, "system": SYSTEM, "prompt": prompt, "format": SCHEMA, "stream": True,
                       "options": {"temperature": 0.2, "num_ctx": 4096, "num_predict": 900}}).encode()
    req = urllib.request.Request(f"{host}/api/generate", data=body, headers={"Content-Type": "application/json"})
    text, stats, t0 = "", {}, time.time()
    with urllib.request.urlopen(req, timeout=1800) as resp:
        for raw in resp:
            if not raw.strip():
                continue
            msg = json.loads(raw)
            if "error" in msg:
                raise RuntimeError(msg["error"])
            text += msg.get("response", "")
            if progress:
                progress("writing", len(re.findall(r"[A-Za-z]{2,}", text)))
            if msg.get("done"):
                n, d = msg.get("eval_count", 0), msg.get("eval_duration", 0) / 1e9
                stats = {"tokens": n, "tokens_per_s": round(n / d, 1) if d else None,
                         "seconds": round(time.time() - t0, 1)}
    return json.loads(text), stats


def ollama_has(model, host=OLLAMA):
    try:
        with urllib.request.urlopen(f"{host}/api/tags", timeout=3) as r:
            names = [m["name"] for m in json.load(r).get("models", [])]
    except (urllib.error.URLError, OSError, ValueError):
        return False, "Ollama is not running on this laptop (open the Ollama app, or run `ollama serve`)"
    if model in names or f"{model}:latest" in names:
        return True, ""
    return False, f"model {model} is not installed (run `ollama pull {model}`)"


# ---------------------------------------------------------------- fact check (code checks the model)
def allowed_tokens(facts, chk):
    times, gs, trees = set(), set(), set()
    for r in chk["records"]:
        times.update({r["dt"][11:16], r["dt"][11:19]})
        if r["name"] in ("DROP", "SHOCK"):
            gs.add(f"{r['value'] / 10:.1f}")
        if r["name"] == "TAMPER":
            trees.add(str(r["value"]))
    for x in facts["incidents"]:            # durations computed by code are facts too
        times.update(re.findall(r"\b\d{2}:\d{2}\b", x["context"]))
    for sq in facts["sequences"]:
        times.update(re.findall(r"\b\d{2}:\d{2}\b", sq))
    level, _ = risk_level(facts)
    return times, gs, trees, facts["type_counts"], level, facts["integrity_ok"]


PLURALS = {"TAMPER": r"(forced[- ]entry|intrusion|break-in|tamper\w*) attempts|break-ins",
           "OPEN !!": r"unauthori[sz]ed (lid )?(openings|accesses|access events|entries)|openings without authori[sz]ation",
           "BADGE ?": r"unknown badges", "DROP": r"\bdrops\b", "FLIPPED": r"\bflips\b",
           "BOOT": r"power (cuts|interruptions|outages)"}


CONTRADICTIONS = {
    "risk": r"(investigation|inspection|action|check)s? (is|are) not (required|needed|necessary)"
            r"|no (further )?(action|investigation|inspection)s? (is |are )?(required|needed|necessary)",
    "integrity": r"(manipulat|tamper|alter|forg|edit)\w* (of |with )?(the )?(log|records)\b"
                 r"|(log|records) (may have been|was|were|could have been|might have been) (manipulat|tamper|alter|forg|edit)",
}


def unsupported(sentence, allowed):
    times, gs, trees, counts, level, integrity_ok = allowed
    bad = [t for t in re.findall(r"\b\d{1,2}:\d{2}(?::\d{2})?\b", sentence) if t.zfill(5) not in times]
    bad += [g + " g" for g in re.findall(r"(\d+(?:\.\d+)?)\s?g\b", sentence) if f"{float(g):.1f}" not in gs]
    bad += [n + "/10" for n in re.findall(r"\b(\d+)\s?/\s?10\b", sentence) if n not in trees]
    low = sentence.lower()
    for kind, pattern in PLURALS.items():               # "multiple unauthorised openings" when there was one
        m = re.search(pattern, low)
        if m and counts.get(kind, 0) < 2:
            bad.append(f"'{m.group(0)}' (the log has {counts.get(kind, 0)})")
    if re.search(r"\b(repeated|multiple|several|numerous)\s+(instances|attempts|incidents of)", low) \
            and max([counts.get(k, 0) for k in SECURITY] + [0]) < 2:
        bad.append("'repeated/multiple' (each security event happened once)")
    if level != "LOW" and re.search(CONTRADICTIONS["risk"], low):    # the verdict belongs to the rules
        bad.append(f"says no action is needed, but the risk level is {level}")
    if integrity_ok and re.search(CONTRADICTIONS["integrity"], low):
        bad.append("suggests the log was manipulated, but its integrity was verified")
    return bad


def readable_ids(text, facts):
    """'(I4, I6)' or 'incident I3' mean nothing to an inspector: replace ids by times."""
    when = {x["id"]: x["time"][11:16] for x in facts["incidents"]}
    text = re.sub(r"\s*\((?:\s*(?:I\d+|linked sequence)\s*,?)+\)", "", text, flags=re.I)
    text = re.sub(r"\b(?:incident\s+)?(I\d+)\b", lambda m: f"the {when[m.group(1)]} incident" if m.group(1) in when
                  else m.group(0), text)
    return text


def clean(text, allowed, rejected, where):
    kept = []
    for s in re.split(r"(?<=[.!?])\s+", text.strip()):
        bad = unsupported(s, allowed)
        if bad:
            rejected.append({"where": where, "sentence": s, "unsupported": bad})
        elif s:
            kept.append(s)
    return " ".join(kept)


def fact_check(draft, facts, chk):
    allowed, rejected = allowed_tokens(facts, chk), []
    if not isinstance(draft, dict):
        draft = {}
    ids_to_time = lambda t: readable_ids(str(t), facts)          # noqa: E731
    ids = {x["id"] for x in facts["incidents"]}
    out = {"executive_summary": clean(ids_to_time(draft.get("executive_summary", "")), allowed, rejected, "summary"),
           "assessments": {}, "recommendations": []}
    for a in draft.get("assessments", []):
        key = str(a.get("id", "")).strip().upper()
        key = "I" + key if key.isdigit() else key          # small models sometimes write "1" for "I1"
        if key in ids:
            out["assessments"][key] = clean(ids_to_time(a.get("assessment", "")), allowed, rejected, key)
    for rec in draft.get("recommendations", [])[:4]:
        c = clean(ids_to_time(rec), allowed, rejected, "recommendation")
        if c:
            out["recommendations"].append(c)
    return out, rejected


# ---------------------------------------------------------------- required actions (inspection rules, code)
def required_actions(facts):
    """Checks the rules impose for each incident type. The model may add suggestions, never remove these."""
    acts, first = [], {}
    for x in facts["incidents"]:
        first.setdefault(x["event"], x)
    if not facts["integrity_ok"]:
        acts.append("Treat this log as unreliable and report the integrity failure.")
    if "OPEN !!" in first:
        acts.append("Count the contents against the packing list and record any shortfall.")
    if "TAMPER" in first:
        acts.append("Photograph and examine the lid, hinges and seals for tool marks before any further handling.")
    if "BADGE ?" in first:
        x = first["BADGE ?"]
        acts.append(f"Identify who held the unregistered badge ({x['detail']}) presented at {x['time'][11:16]}.")
    if "BOOT" in first:
        acts.append(f"Obtain the custody and escort records for the period the crate was unpowered "
                    f"(before the restart at {first['BOOT']['time'][11:16]}).")
    if "CLEARED" in first:
        acts.append("Find out who erased the log; treat the history before the erasure as unverified.")
    if "DROP" in first or "FLIPPED" in first:
        acts.append("Check the contents for impact damage or displacement.")
    return acts


def merge_actions(required, suggested, limit=6):
    """Required actions first, then the model's suggestions that do not repeat them."""
    words = lambda t: set(re.findall(r"[a-z]{4,}", t.lower()))          # noqa: E731
    out = list(required)
    for s in suggested:
        if len(out) >= limit:
            break
        if all(len(words(s) & words(o)) / max(1, len(words(s) | words(o))) < 0.4 for o in out):
            out.append(s)
    return out


# ---------------------------------------------------------------- fallback without a model
TEMPLATE_ASSESSMENT = {
    "OPEN !!": "The lid was opened without a valid badge in the preceding 30 seconds; the contents should be counted against the packing list.",
    "TAMPER": "The motion pattern matches forced entry; check the lid, hinges and seals for marks.",
    "BADGE ?": "A badge not registered for this crate was presented; identify who held it.",
    "DROP": "The crate fell and hit the ground; check the contents for impact damage.",
    "FLIPPED": "The crate was turned upside down; check that the contents did not shift.",
    "BOOT": "Monitoring stopped while the power was off; events during that period are unknown.",
    "CLEARED": "Earlier history was erased; treat the shipment as unverified before this point.",
    "SHOCK": "A strong shock was recorded; check the contents for damage.",
}


def template_draft(facts, level, reasons):
    n = len(facts["incidents"])
    summary = (f"The log of crate {facts['crate_id']} covers {facts['start']} to {facts['end']} and "
               f"{'passed' if facts['integrity_ok'] else 'FAILED'} its integrity check. "
               f"{n} incident{'s' if n != 1 else ''} recorded. Risk level {level}: {'; '.join(reasons)}.")
    recs = (["Open the crate and count the contents against the packing list.",
             "Photograph the lid, seals and hinges before any further handling."] if level == "HIGH" else
            ["Inspect the contents for transport damage."] if level == "MEDIUM" else
            ["No action needed beyond the routine check."])
    if not facts["integrity_ok"]:
        recs.insert(0, "Treat this log as unreliable and report the integrity failure.")
    return {"executive_summary": summary,
            "assessments": {x["id"]: TEMPLATE_ASSESSMENT.get(x["event"], "") for x in facts["incidents"]},
            "recommendations": recs}


# ---------------------------------------------------------------- Word document
def write_docx(rep, path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    colors = {"HIGH": RGBColor(0xC0, 0x1F, 0x2B), "MEDIUM": RGBColor(0x9A, 0x67, 0x00), "LOW": RGBColor(0x1A, 0x7F, 0x37)}
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)

    def shade(cell, hex_color):
        tc = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
        tc.append(shd)

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(h)
            run.bold, run.font.size = True, Pt(9.5)
            shade(c, "E8EBED")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(v))
                run.font.size = Pt(9.5)
        t.autofit = False
        for i, w in enumerate(widths):                 # Word reads the cells, LibreOffice the grid
            t.columns[i].width = Cm(w)
            for row in t.rows:
                row.cells[i].width = Cm(w)
        return t

    title = doc.add_heading("WARDEN inspection report", level=0)
    title.runs[0].font.size = Pt(24)
    p = doc.add_paragraph()
    p.add_run(f"Crate {rep['crate_id']}   ·   log period {rep['period'][0]} to {rep['period'][1]} (crate clock)   ·   "
              f"generated {rep['generated']}").font.size = Pt(9)

    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    for c, (k, v, col) in zip(t.rows[0].cells + t.rows[1].cells, [
            ("RISK LEVEL", rep["risk"]["level"], colors[rep["risk"]["level"]]),
            ("LOG INTEGRITY", "VERIFIED" if rep["integrity"]["ok"] else "FAILED",
             colors["LOW"] if rep["integrity"]["ok"] else colors["HIGH"]),
            ("WHY", "; ".join(rep["risk"]["reasons"]), None),
            ("RECORDS", f"{rep['integrity']['records']} records, hash chain "
                        + ("unbroken" if rep["integrity"]["ok"] else "broken: " + "; ".join(rep["integrity"]["issues"])), None)]):
        c.text = ""
        r1 = c.paragraphs[0].add_run(k + "\n")
        r1.font.size, r1.bold = Pt(8), True
        r2 = c.paragraphs[0].add_run(v)
        r2.font.size, r2.bold = (Pt(16), True) if col else (Pt(9.5), False)
        if col:
            r2.font.color.rgb = col
        shade(c, "F6F7F8")

    ai = rep["meta"]["engine"] == "llm"
    tag = f"Drafted by {rep['meta']['model']} running offline on this laptop, checked against the log by code." if ai \
        else "Template text (no local language model available)."
    doc.add_heading("1. Summary", level=1)
    doc.add_paragraph(rep["summary"])
    doc.add_paragraph().add_run(tag).italic = True

    doc.add_heading("2. Incidents", level=1)
    if rep["incidents"]:
        table(["Time", "Event", "Facts from the log", "Assessment"],
              [(x["time"], x["label"], "; ".join(v for v in (x["detail"], x["context"]) if v), x["assessment"] or "-")
               for x in rep["incidents"]], [2.6, 3.2, 5.0, 5.8])
    else:
        doc.add_paragraph("No incident recorded.")
    for sq in rep.get("sequences", []):
        p = doc.add_paragraph()
        p.add_run("Linked sequence: ").bold = True
        p.add_run(sq)

    doc.add_heading("3. Recommended actions", level=1)
    for r in rep["recommendations"]:
        doc.add_paragraph(r, style="List Number")
    n = rep.get("required_actions", 0)
    if n and len(rep["recommendations"]) > n:
        note = f"Actions 1 to {n} are required by the inspection rules for these incident types; the others are suggestions."
    elif n:
        note = "These actions are required by the inspection rules for these incident types."
    else:
        note = ""
    if note:
        doc.add_paragraph().add_run(note).italic = True

    doc.add_heading("4. Full event log", level=1)
    table(["#", "Date and time", "Event", "Value", "Chain"],
          [(r["seq"], r["dt"], LABELS.get(r["name"], r["name"]), detail(r), r["status"]) for r in rep["timeline"]],
          [1.0, 3.6, 6.0, 3.6, 2.4])

    doc.add_heading("5. Method and limits", level=1)
    method = [
        "Events are detected on the crate: drops and flips by physical rules, forced entry by a random forest running "
        "on the crate's microcontroller, openings by the light sensor and RFID badges. Each record is chained to the "
        "previous one with a CRC-16 fingerprint; the chain was re-verified on this laptop before writing this report.",
        "The risk level and the required actions are decided by fixed rules, not by the language model. The model only "
        "writes the summary, the assessments and optional extra suggestions; every time, figure and count it writes is "
        "checked against the verified log, sentences contradicting the verdict are rejected, and "
        f"{len(rep['meta']['rejected'])} sentence(s) that could not be matched were removed.",
        "CRC-16 detects edits by someone who does not know the algorithm; a fielded version would use a keyed MAC "
        "with the key held in a secure element. Times come from the crate's own clock.",
    ]
    if ai and rep["meta"].get("tokens_per_s"):
        method.append(f"Model: {rep['meta']['model']} via Ollama, {rep['meta']['seconds']} s, "
                      f"{rep['meta']['tokens_per_s']} tokens/s, no network connection required.")
    for m in method:
        doc.add_paragraph(m).runs[0].font.size = Pt(9)

    doc.add_paragraph()
    doc.add_paragraph("Inspector: ______________________     Date: ____________     Signature: ______________________")
    doc.save(path)


# ---------------------------------------------------------------- the whole thing
def generate(lines, model=DEFAULT_MODEL, crate_id="WARDEN-01", use_llm=True, progress=None, out_dir=REPORTS):
    progress = progress or (lambda *a: None)
    progress("verifying", 0)
    chk = V.check(lines)
    facts = build_facts(chk, crate_id)
    level, reasons = risk_level(facts)
    meta = {"engine": "template", "model": None, "rejected": [], "note": ""}

    draft = None
    if use_llm and chk["records"]:
        ok, why = ollama_has(model)
        if ok:
            progress("writing", 0)
            try:
                raw, stats = ask_llm(TASK.format(facts=facts_text(facts, level, reasons)), model, progress)
                progress("checking", 0)
                draft, rejected = fact_check(raw, facts, chk)
                meta.update(engine="llm", model=model, rejected=rejected, **stats)
                if not draft["executive_summary"]:       # everything rejected: do not ship an empty summary
                    draft["executive_summary"] = template_draft(facts, level, reasons)["executive_summary"]
            except (urllib.error.URLError, OSError, ValueError, RuntimeError) as e:
                meta["note"] = f"local model failed ({e}); template used"
        else:
            meta["note"] = why
    if draft is None:
        draft = template_draft(facts, level, reasons)
    fallback = template_draft(facts, level, reasons)
    required = required_actions(facts)
    actions = merge_actions(required, draft["recommendations"] or ([] if required else fallback["recommendations"]))

    rep = {
        "crate_id": crate_id, "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "period": [facts.get("start") or "-", facts.get("end") or "-"],
        "integrity": {"ok": chk["ok"], "issues": chk["issues"], "records": len(chk["records"])},
        "risk": {"level": level, "reasons": reasons},
        "summary": draft["executive_summary"],
        "incidents": [dict(x, assessment=draft["assessments"].get(x["id"]) or fallback["assessments"].get(x["id"], ""))
                      for x in facts["incidents"]],
        "recommendations": actions,
        "required_actions": len(required),
        "sequences": facts.get("sequences", []),
        "timeline": chk["records"],
        "meta": meta,
    }
    progress("document", 0)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    name = f"{crate_id}_{stamp}"
    write_docx(rep, out_dir / f"{name}.docx")
    rep["docx"] = f"{name}.docx"
    (out_dir / f"{name}.json").write_text(json.dumps(rep, indent=1))
    progress("done", 0)
    return rep


# ---------------------------------------------------------------- demo log (valid chain)
def demo_lines():
    """A plausible three-day shipment with a night-time break-in, a drop and a power cut."""
    base = calendar.timegm(datetime(2026, 11, 3, 7, 50).timetuple())   # crate clock: local time stored as Unix
    events = [(0, "BOOT", 0), (600, "AUTH", 0), (612, "OPEN OK", 190), (1500, "CLOSED", 9),
              (14520, "DROP", 34), (31800, "FLIPPED", 0), (32100, "UPRIGHT", 0),
              (66840, "TAMPER", 9), (66905, "OPEN !!", 181), (67560, "CLOSED", 11),
              (69300, "BOOT", 0), (129600, "BADGE ?", 0x3A), (175200, "INSPECT", 0)]
    lines, chain = ["#L,seq,unix,datetime,event,value,chain"], 0xFFFF
    for seq, (dt, name, value) in enumerate(events):
        unix = base + dt
        chain = V.crc16(chain, unix.to_bytes(4, "little") + bytes([V.EVENT_CODES[name], value]))
        stamp = (datetime(1970, 1, 1) + timedelta(seconds=unix)).strftime("%Y-%m-%d %H:%M:%S")
        lines.append(f"#L,{seq},{unix},{stamp},{name},{value},{chain:X}")
    lines.append(f"#L,END,{len(events)},{chain:X}")
    return lines


def main():
    ap = argparse.ArgumentParser(description="Draft a WARDEN inspection report with a local language model.")
    src = ap.add_mutually_exclusive_group()
    src.add_argument("--file", help="a saved dump (verify_log.py --save)")
    src.add_argument("--demo", action="store_true", help="use a built-in sample shipment")
    ap.add_argument("--port", help="serial port of the crate (auto-detected if omitted)")
    ap.add_argument("--model", default=DEFAULT_MODEL, help=f"Ollama model (default {DEFAULT_MODEL})")
    ap.add_argument("--no-llm", action="store_true", help="template text only")
    ap.add_argument("--crate", default="WARDEN-01", help="crate identifier printed on the report")
    args = ap.parse_args()

    if args.demo:
        lines = demo_lines()
    elif args.file:
        lines = [l.strip() for l in open(args.file) if l.startswith("#L,")]
    else:
        lines = V.read_from_board(args.port)

    last = {"stage": None}

    def progress(stage, n):
        if stage == "done":
            return
        if stage != last["stage"]:
            print({"verifying": "Verifying the log ...", "writing": f"Local model {args.model} writing ...",
                   "checking": "\nChecking every time and figure against the log ...",
                   "document": "Writing the Word document ...", "done": ""}[stage], end="", flush=True)
            last["stage"] = stage
            if stage != "writing":
                print()
        elif stage == "writing":
            print(f"\rLocal model {args.model} writing ... {n} words", end="", flush=True)

    rep = generate(lines, args.model, args.crate, not args.no_llm, progress)
    m = rep["meta"]
    print(f"\nRisk level: {rep['risk']['level']} ({'; '.join(rep['risk']['reasons'])})")
    print(f"Integrity: {'VERIFIED' if rep['integrity']['ok'] else 'FAILED'}")
    print(f"\n{rep['summary']}\n")
    if m["engine"] == "llm":
        speed = f", {m['tokens_per_s']} tokens/s" if m.get("tokens_per_s") else ""
        print(f"Written by {m['model']} in {m.get('seconds')} s{speed}; "
              f"{len(m['rejected'])} unsupported sentence(s) removed.")
        for r in m["rejected"]:
            print(f"  removed ({r['where']}): {r['sentence']}  [not in the log: {', '.join(r['unsupported'])}]")
    else:
        print(f"Template text used: {m['note'] or 'requested'}")
    print(f"\nReport: {REPORTS / rep['docx']}")


if __name__ == "__main__":
    main()
